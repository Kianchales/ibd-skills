#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""overview_to_docx.py — 复核总览报告 MD → Word 转换器

为什么存在（2026-09-18 用户裁定）：
    总览报告交付 = **MD 版 + Word 版双格式**。MD 供程序/检索，Word 供批阅流转。
    annotate_docx.py / annotate_pdf.py 生成 `_批注总览.md` 后调用本脚本同产 `.docx`。

转换范围（只认本包 write_overview 产出的总览子集，不追求通用 md 渲染）：
    `# `/`## ` 标题、`>` 引导行、`|` 表格（首行表头加粗）、`- `/`  - ` 两级列表、
    `**加粗**` 行内加粗、`` `代码` `` 去反引号、空行分段。

用法：
    python overview_to_docx.py <总览.md> [--out <总览.docx>]
    （缺省 --out = 同名 .docx；被 annotate 脚本内自动调用时传 md 路径即可）
退出码：0 = 成功；1 = 失败；2 = 用法错误。
"""
import argparse
import os
import re
import sys

try:
    from docx import Document
except ImportError:
    sys.exit("缺少 python-docx：annotate docx 链路的既有依赖，请先安装")


def add_runs(p, text):
    """渲染行内格式：**加粗** 拆 run 加粗；`代码` 去反引号。"""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            r.bold = True
        else:
            p.add_run(part.replace("`", ""))


def _flush_table(doc, rows):
    """rows = 已按 | 切好单元格的行列表（不含分隔行）。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.paragraphs[0].text = ""
            txt = row[j] if j < len(row) else ""
            add_runs(cell.paragraphs[0], txt)
            if i == 0:  # 首行表头加粗
                for r in cell.paragraphs[0].runs:
                    r.bold = True


def convert(md_path, out_path=None):
    if not os.path.isfile(md_path):
        print(f"[ERROR] 总览 md 不存在: {md_path}")
        return 1
    out = out_path or os.path.splitext(md_path)[0] + ".docx"
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    table_buf = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        stripped = ln.strip()

        # ── 表格块（连续 | 行；第 2 行为分隔行则跳过） ──
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            # 分隔行（---）不入表
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table_buf.append(cells)
            i += 1
            continue
        if table_buf:
            _flush_table(doc, table_buf)
            table_buf = []

        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].replace("`", ""), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].replace("`", ""), level=1)
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = None
            add_runs(p, stripped.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
        elif ln.startswith("  - "):
            p = doc.add_paragraph(style="List Bullet 2")
            add_runs(p, stripped[2:])
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)
        i += 1
    if table_buf:
        _flush_table(doc, table_buf)

    doc.save(out)
    print(f"[OK] 总览 Word 版已生成: {out}")
    return 0


def main():
    ap = argparse.ArgumentParser(description="复核总览报告 MD → Word（双格式交付）")
    ap.add_argument("md", help="总览 md 路径（write_overview 产出）")
    ap.add_argument("--out", help="输出 docx 路径（默认同名 .docx）")
    args = ap.parse_args()
    sys.exit(convert(args.md, args.out))


if __name__ == "__main__":
    main()
