#!/usr/bin/env python3
"""review_report_to_xlsx.py — 复核报告（Excel）生成器

为什么存在（2026-10-07 用户裁定）：
    复核报告的交付载体由 **Word 改为 Excel**——「Excel 可以根据问题性质分成不同的 sheet，
    比 Word 报告好」。此前该报告的**文件级骨架从未进规范**（`ibd-doc-review` delivery.md
    §八 的「收尾四段式」是**对话回复**的话术骨架，不是报告文件骨架），本脚本与
    `ibd-doc-review/references/delivery.md` **§八之二**配套：**规范定形态，本脚本落形态**。

输入：
    --issues <问题清单.json>   复核问题清单（结构契约 ibd-doc-review/references/interface.md §3）
    --meta  <报告元信息.json>  **可选**——报告头的叙述性内容（复核对象/基准/方式/**复核声明**/
                               范围与方法/已核查通过/质量证据/行动项）；缺省时只出明细页 ＋ 自动统计。
                               模板随包：scripts/report-meta.example.json
输出：
    Excel 工作簿。G3 命名规范：`<原文>_<YYYYMMDD>_v<N>_复核报告.xlsx`（正式交付请用 --out 显式传入）。

工作簿结构（规格单一事实源 = delivery.md §八之二）：
    ① 封面与汇总   报告头 ＋ **复核声明** ＋ 结论 ＋ 复核范围与方法 ＋ 统计 ＋ 已核查通过
                   ＋ 质量证据 ＋ 行动项（小节序号按**实际渲染顺序**自增，缺节不跳号；
                   **仅主标题一行居中**，其余标题一律左对齐；**所有行统一跨 A:C 合并**，
                   列宽 A:B:C 唯一一套 —— 封面上下等宽）
    ② 按性质分页   一张工作表一个「问题性质」（--split 可换轴；nature 缺省回退 type 大类）
    ③ 全部         不拆分的完整台账

明细列：编号｜严重度｜性质｜域｜行号｜原文片段｜问题标题｜问题描述｜建议（＋定稿文本/作者，有则加列）

用法：
    python review_report_to_xlsx.py --issues issues.json [--meta meta.json]
           [--out 报告.xlsx] [--split nature] [--title "…"] [--date 20261007] [--json]
退出码：0 = 成功；1 = 失败；2 = 用法/清单错误。

依赖：openpyxl（**惰性导入**——仅本功能需要，本包其它链路不依赖）
编号同源：issue_numbering.py（与批注注入器共用，保证「报告编号 ↔ 批注编号 一一对应」）
"""
import argparse
import copy
import datetime
import json
import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# ── 与注入器共用的编号单一事实源（零三方依赖）──
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from issue_numbering import derive_code, full_label  # noqa: E402
from validate_issues import validate_issues  # noqa: E402

# ── 展示常量 ──
SEV_ORDER = {"高": 0, "中": 1, "低": 2}
SEV_LABEL = {"高": "必改（高）", "中": "应改（中）", "低": "建议改（低）"}
# 性质标签的排序优先级：**只看前导圈号**（①–⑩），不比对完整标签文字——标签由项目自定
# （如「③自身逻辑」vs「③回复自身逻辑」），圈号才是跨项目稳定的序。无圈号者排末位。
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"
SPLIT_CHOICES = ("nature", "type", "code", "sev")

SHEET_COVER = "封面与汇总"
SHEET_ALL = "全部"

# 封面统一跨度：**所有行一律合并/铺设到 C 列**（用户裁定 2026-10-07）——
# 此前多数行只跨 A:B，而「域／覆盖范围／条数」统计表天然占 A:B:C，导致封面**上下不等宽**。
COVER_LAST_COL = 3
COVER_WIDTHS = {"A": 20, "B": 78, "C": 20}   # 合计 118，全封面唯一一套列宽

# 封面小节序号（「一、」「二、」…）——按渲染顺序自增，某节缺省时不跳号
CN_NUM = "一二三四五六七八九十"

# 复核声明的分栏（键 → 小标题），对应《本轮检查项声明》
# （ibd-doc-review check-scope.md §六）；「声明与报告头一致」＝交付纪律，见 delivery.md §八。
DECL_SECTIONS = (("checked", "检查（本次实际要跑的项）"),
                 ("not_checked", "不查（点名 ＋ 原因）"),
                 ("special", "特别专项（触发即生效）"),
                 ("gates", "交付前将跑"))

# 严重度着色：高＝红系｜中＝橙系｜低＝灰系（遵中文语境「红＝最需处理」）
SEV_STYLE = {"高": ("FFC7CE", "9C0006"), "中": ("FFEB9C", "9C6500"), "低": ("EDEDED", "595959")}
HEAD_FILL = "D9E1F2"
HEAD_FONT = "1F3864"
SECTION_FONT = "1F3864"
BORDER_COLOR = "BFBFBF"

# 台账列（键 → 表头），顺序即列序；末两列按需追加
BASE_COLS = [("full", "编号"), ("sev", "严重度"), ("nature", "性质"), ("code", "域"),
             ("line", "行号"), ("anchor", "原文片段"), ("title", "问题标题"),
             ("desc", "问题描述"), ("advice", "建议")]
OPT_COLS = [("rev", "定稿文本"), ("author", "作者")]
COL_WIDTH = {"编号": 10, "严重度": 11, "性质": 16, "域": 6, "行号": 13, "原文片段": 52,
             "问题标题": 34, "问题描述": 52, "建议": 40, "定稿文本": 46, "作者": 10}
LONG_COLS = {"原文片段", "问题标题", "问题描述", "建议", "定稿文本"}


def _load_openpyxl():
    """惰性导入：openpyxl 只在真正生成报告时需要，不拖累本包其它链路。"""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover
        sys.exit("[ERROR] 缺少 openpyxl——复核报告生成器需要它（pip install openpyxl）；"
                 "本包其余链路不受影响。")
    return openpyxl, Alignment, Border, Font, PatternFill, Side, get_column_letter


# ───────────────────────── 领域逻辑（无 Excel 依赖，便于单测） ─────────────────────────

def nature_of(it):
    """条目的「性质」：清单 `nature` 字段优先；缺省回退 `type` 的大类（数据/表述/披露/合规）。"""
    n = str(it.get("nature") or "").strip()
    if n:
        return n
    t = str(it.get("type") or "").strip()
    if "·" in t:
        return t.split("·", 1)[0]
    return t or "未分类"


def number_entries(issues):
    """按注入器同一规则编号（issue_numbering）——**不改动入参**，返回新列表。

    每组前缀独立自增，顺序 ＝ 清单顺序（编号一次性分配，与批注/总览一一对应）。
    """
    seen, out = {}, []
    for it in issues:
        code = derive_code(it)
        seen[code] = seen.get(code, 0) + 1
        row = dict(it)
        row["code"] = code
        row["full"] = full_label(code, seen[code])
        row["nature"] = nature_of(it)
        row["sev"] = str(it.get("sev") or "")
        out.append(row)
    return out


def _nature_rank(name):
    """性质标签排序键：有前导圈号①–⑩按圈号；无则排到末位（同档再按条数降、名称升）。"""
    return CIRCLED.index(name[0]) if name and name[0] in CIRCLED else len(CIRCLED)


def group_entries(rows, axis):
    """按轴分组，返回 [(组名, [条目…])]，组序可预期（推荐序 > 条数 > 名称）。"""
    key = {"nature": lambda r: r["nature"],
           "type": lambda r: (r["type"].split("·", 1)[0] if "·" in r["type"] else r["type"]) or "未分类",
           "code": lambda r: r["code"],
           "sev": lambda r: r["sev"] or "未定"}[axis]
    buckets = {}
    for r in rows:
        buckets.setdefault(key(r), []).append(r)
    if axis == "sev":
        names = sorted(buckets, key=lambda k: (SEV_ORDER.get(k, 9), k))
    elif axis == "nature":
        names = sorted(buckets, key=lambda k: (_nature_rank(k), -len(buckets[k]), k))
    else:
        names = sorted(buckets, key=lambda k: (len(k) != 1, k))
    return [(n, buckets[n]) for n in names]


def sort_rows(rows):
    """严重度升序（高→低），同档保持清单原序。"""
    return sorted(rows, key=lambda r: SEV_ORDER.get(r["sev"], 9))


def counts_by(rows, axis):
    if axis == "sev":
        return [(SEV_LABEL.get(k, k), sum(1 for r in rows if r["sev"] == k))
                for k in ("高", "中", "低") if any(r["sev"] == k for r in rows)]
    if axis == "code":
        return [(n, len(g)) for n, g in group_entries(rows, "code")]
    if axis == "type":
        return [(n, len(g)) for n, g in group_entries(rows, "type")]
    return [(n, len(g)) for n, g in group_entries(rows, "nature")]


def _sanitize_sheet_name(name, used):
    """Excel 工作表名净化：去 []:*?/\\ 、截 31、去首尾引号，冲突时加序号。"""
    s = re.sub(r"[\[\]:*?/\\]", "", str(name)).strip().strip("'").strip()
    s = s[:31] or "未命名"
    base, i = s, 2
    while s in used or s == SHEET_COVER:
        suffix = f"({i})"
        s = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(s)
    return s


# ───────────────────────── Excel 渲染 ─────────────────────────

def _write_table(ws, r, header, rows, widths, styles, start_col=1, extend_last_to=None):
    """写一张带表头的表；返回下一空行号。

    `extend_last_to`：把**末列右延并合并**到指定列（封面统计表统一到 C 列用，
    保证封面「上下等宽」——见 COVER_LAST_COL）。
    """
    Alignment, Border, Font, PatternFill, Side, _ = styles
    thin = Side(style="thin", color=BORDER_COLOR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill, head_font = HEAD_FILL, HEAD_FONT
    n = len(header)
    last_native = start_col + n - 1
    end_col = max(last_native, extend_last_to or 0)

    def _extend(rr):
        """末列右延到 end_col：先给补位格描边、再合并（合并区外框取各边所属格）。"""
        if end_col <= last_native:
            return
        for j in range(last_native + 1, end_col + 1):
            ws.cell(row=rr, column=j).border = box
        ws.merge_cells(start_row=rr, start_column=last_native, end_row=rr, end_column=end_col)

    for j, h in enumerate(header):
        c = ws.cell(row=r, column=start_col + j, value=h)
        c.fill = PatternFill("solid", fgColor=head_fill)
        c.font = Font(bold=True, color=head_font)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    _extend(r)
    r += 1
    sev_idx = header.index("严重度") if "严重度" in header else -1
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=start_col + j, value=v)
            c.border = box
            c.alignment = Alignment(vertical="top",
                                    wrap_text=header[j] in LONG_COLS,
                                    horizontal="center" if header[j] in ("编号", "严重度", "性质", "域", "行号") else "left")
            if j == sev_idx and str(v) in SEV_STYLE:
                fill, font = SEV_STYLE[str(v)]
                c.fill = PatternFill("solid", fgColor=fill)
                c.font = Font(bold=True, color=font)
        _extend(r)
        r += 1
    return r


def _write_kv(ws, r, pairs, styles, key_w=None):
    """标签格（A）＋ 值格（**合并 B:C**）——值格右延到 COVER_LAST_COL，与全封面等宽。"""
    Alignment, Border, Font, PatternFill, Side, _ = styles
    thin = Side(style="thin", color=BORDER_COLOR)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for k, v in pairs:
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True)
        a.fill = PatternFill("solid", fgColor="F2F2F2")
        a.border = box
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row=r, column=2, value=v)
        b.border = box
        b.alignment = Alignment(vertical="top", wrap_text=True)
        for j in range(3, COVER_LAST_COL + 1):
            ws.cell(row=r, column=j).border = box
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=COVER_LAST_COL)
        r += 1
    return r


def _write_center(ws, r, text, styles, bold=False, size=11, color=None):
    """跨 **A:C** 合并的**居中**单行——**仅封面主标题（第 1 行）使用**。

    用户裁定（2026-10-07）：**只第一行居中**，其余标题一律正常居左——故本函数
    不再用于副标题，也不用于小节标题（见 `_write_section`）。
    """
    Alignment, _, Font, _, _, _ = styles
    c = ws.cell(row=r, column=1, value=text)
    opts = {"bold": bold, "size": size}
    if color:
        opts["color"] = color
    c.font = Font(**opts)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COVER_LAST_COL)
    return r + 1


def _write_section(ws, r, title, styles):
    """小节标题：跨 **A:C** 合并 ＋ **左对齐**（用户裁定：主标题之外一律正常居左；
    跨度仍与全封面统一，保证上下等宽）。"""
    Alignment, _, Font, _, _, _ = styles
    c = ws.cell(row=r, column=1, value=title)
    c.font = Font(bold=True, size=12, color=SECTION_FONT)
    c.alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COVER_LAST_COL)
    return r + 1


def _write_lines(ws, r, lines, styles, bullet=True):
    Alignment, _, _, _, _, _ = styles
    for ln in lines:
        c = ws.cell(row=r, column=1, value=("· " if bullet else "") + str(ln))
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COVER_LAST_COL)
        r += 1
    return r


def _write_declaration(ws, r, decl, styles):
    """**复核声明**块（2026-10-07 用户裁定：声明要随报告走）。

    报告是一份独立文件、会脱离对话单独流转，故「动手前声明什么、报告头就写什么」
    （`ibd-doc-review` delivery.md §八 落款义务 ⟷ check-scope.md §六 声明单）。
    两种形态都收：
      - dict：{depth, checked[], not_checked[], special[], gates}（结构化，推荐）
      - list／str：逐行文字（简单形态，无分栏）
    """
    Alignment, _, Font, _, _, _ = styles
    if isinstance(decl, str):
        decl = [decl]
    if isinstance(decl, list):
        return _write_lines(ws, r, decl, styles)
    depth = str(decl.get("depth") or "").strip()
    if depth:
        c = ws.cell(row=r, column=1, value=f"复核深度：{depth}")
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COVER_LAST_COL)
        r += 1
    for key, label in DECL_SECTIONS:
        val = decl.get(key)
        if not val:
            continue
        c = ws.cell(row=r, column=1, value="▍" + label)
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COVER_LAST_COL)
        r += 1
        r = _write_lines(ws, r, val if isinstance(val, list) else [val], styles)
    return r


class _SectionNumber:
    """封面小节序号自增器——**按实际渲染顺序**编号，某节缺省时不出现跳号。"""

    def __init__(self):
        self._i = 0

    def __call__(self, title):
        self._i += 1
        n = CN_NUM[self._i - 1] if self._i <= len(CN_NUM) else str(self._i)
        return f"{n}、{title}"


def _detail_sheet(wb, name, rows, cols, styles, first=False):
    Alignment, _, Font, _, _, get_column_letter = styles
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    header = [h for _, h in cols]
    data = [[r.get(k, "") for k, _ in cols] for r in sort_rows(rows)]
    nxt = _write_table(ws, 1, header, data, None, styles)
    for j, (_, h) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = COL_WIDTH.get(h, 14)
    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"
    return ws


def _cover_sheet(wb, meta, rows, split, styles):
    Alignment, _, Font, _, _, get_column_letter = styles
    ws = wb.active
    ws.title = SHEET_COVER
    # 统一三列宽（A:B:C）——全封面唯一一套，任何行都不再只跨 A:B（用户裁定：上下等宽）
    for col, w in COVER_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # 主标题：**仅此一行居中**（用户裁定 2026-10-07）；副标题与其余标题一律左对齐
    title = str(meta.get("title") or "复核报告")
    r = _write_center(ws, 1, title, styles, bold=True, size=14, color=SECTION_FONT)
    total = len(rows)
    hi = sum(1 for x in rows if x["sev"] == "高")
    mid = sum(1 for x in rows if x["sev"] == "中")
    lo = sum(1 for x in rows if x["sev"] == "低")
    r = _write_lines(ws, r,
                     [f"生成日期：{meta.get('date') or datetime.date.today().strftime('%Y-%m-%d')}　｜　"
                      f"本报告编号与《批注版》原文批注一一对应（同源同一份问题清单）。"],
                     styles, bullet=False)
    r += 1

    # 报告头（非编号节）
    r = _write_section(ws, r, "报告头", styles)
    r = _write_kv(ws, r, [
        ("复核对象", meta.get("subject", "—")),
        ("唯一数据基准", meta.get("baseline", "—")),
        ("复核方式", meta.get("method", "—")),
        ("复核结论", meta.get("conclusion")
         or f"共确认 {total} 条问题：必改（高）{hi} ／ 应改（中）{mid} ／ 建议改（低）{lo}"),
    ], styles)
    r += 1

    num = _SectionNumber()

    # 一、复核声明（**声明在先**——动手前声明什么，报告头就写什么，delivery.md §八 落款义务）
    if meta.get("declaration"):
        r = _write_section(ws, r, num("复核声明"), styles)
        r = _write_declaration(ws, r, meta["declaration"], styles)
        r += 1

    # 二、结论
    if meta.get("conclusion_notes"):
        r = _write_section(ws, r, num("结论"), styles)
        r = _write_lines(ws, r, meta["conclusion_notes"], styles)
        r += 1

    # 三、复核范围与方法
    if meta.get("scope") or meta.get("approach"):
        r = _write_section(ws, r, num("复核范围与方法"), styles)
        if meta.get("scope"):
            r = _write_kv(ws, r, [("复核范围", meta["scope"])], styles)
        if meta.get("approach"):
            r = _write_lines(ws, r, meta["approach"], styles)
        r += 1

    # 四、统计（三张表一律铺满 A:C——两张两列表把末列右延到 C，与「域」表等宽）
    r = _write_section(ws, r, num("统计"), styles)
    r = _write_table(ws, r, ["按严重度", "条数"], counts_by(rows, "sev") + [("合计", total)],
                     None, styles, 1, extend_last_to=COVER_LAST_COL)
    r += 1
    r = _write_table(ws, r, ["按性质", "条数"], counts_by(rows, "nature") + [("合计", total)],
                     None, styles, 1, extend_last_to=COVER_LAST_COL)
    r += 1
    scope_map = meta.get("domain_scope") or {}
    dom = sorted(counts_by(rows, "code"), key=lambda kv: kv[0])
    dom_rows = [(k, scope_map.get(k, "—"), v) for k, v in dom] + [("合计", "", total)]
    r = _write_table(ws, r, ["域", "覆盖范围", "条数"], dom_rows, None, styles, 1)
    r += 1
    if split != "nature":
        r = _write_section(ws, r, f"（本报告按「{split}」分工作表；明细见后各页与「{SHEET_ALL}」页）", styles)
        r += 1

    # 五~七
    for sec_title, key in (("已核查通过（关键数据与基准一致）", "verified"),
                           ("质量证据", "quality_evidence"),
                           ("行动项", "actions")):
        if meta.get(key):
            r = _write_section(ws, r, num(sec_title), styles)
            r = _write_lines(ws, r, meta[key], styles)
            r += 1

    # 工作表导航（非编号节）
    r = _write_section(ws, r, "工作表导航", styles)
    names = [SHEET_COVER] + [n for n, _ in group_entries(rows, split)] + [SHEET_ALL]
    r = _write_lines(ws, r, names, styles)
    return ws


# ───────────────────────── 主流程 ─────────────────────────

def build(args, issues, meta):
    openpyxl, Alignment, Border, Font, PatternFill, Side, get_column_letter = _load_openpyxl()
    styles = (Alignment, Border, Font, PatternFill, Side, get_column_letter)

    rows = number_entries(issues)
    cols = list(BASE_COLS) + [(k, h) for k, h in OPT_COLS if any(r.get(k) for r in rows)]
    meta = dict(meta)
    if args.title:
        meta["title"] = args.title
    if args.date:
        meta["date"] = args.date
    else:
        d = str(meta.get("date") or "").strip()
        meta["date"] = d or datetime.date.today().strftime("%Y-%m-%d")

    wb = openpyxl.Workbook()
    _cover_sheet(wb, meta, rows, args.split, styles)

    used = {SHEET_COVER}
    for name, group in group_entries(rows, args.split):
        _detail_sheet(wb, _sanitize_sheet_name(name, used), group, cols, styles)
    _detail_sheet(wb, _sanitize_sheet_name(SHEET_ALL, used), rows, cols, styles)

    out = args.out or os.path.splitext(os.path.abspath(args.issues))[0] + "_复核报告.xlsx"
    wb.save(out)
    return out, rows, cols


def main():
    ap = argparse.ArgumentParser(
        description="复核报告（Excel）生成器（规格：ibd-doc-review references/delivery.md §八之二）")
    ap.add_argument("--issues", required=True, help="复核问题清单 JSON（数组）")
    ap.add_argument("--meta", help="报告元信息 JSON（可选：报告头/范围与方法/已核查/质量证据/行动项）")
    ap.add_argument("--out", help="输出 xlsx（默认 <issues 同名>_复核报告.xlsx；正式交付按 G3 命名显式传入）")
    ap.add_argument("--split", choices=SPLIT_CHOICES, default="nature",
                    help="分工作表轴：nature（默认，问题性质）/ type（12 类大类）/ code（域）/ sev（严重度）")
    ap.add_argument("--title", help="报告标题（覆盖 meta.title）")
    ap.add_argument("--date", help="生成日期（覆盖 meta.date；默认今天）")
    ap.add_argument("--json", action="store_true", help="输出结构化 JSON（供上层消费；不打印人类可读摘要）")
    ap.add_argument("--verbose", action="store_true", help="逐条打印入口校验提示（默认折叠为一行计数）")
    args = ap.parse_args()

    try:
        with open(args.issues, encoding="utf-8") as fh:
            issues = json.load(fh)
    except (OSError, json.JSONDecodeError) as e:
        print(f"[ERROR] 清单读取/解析失败：{e}", file=sys.stderr)
        sys.exit(2)

    # 入口校验（与注入器同源）：结构性坏清单不生成报告
    errs, warns = validate_issues(issues)
    if errs:
        print(f"[ERROR] 清单未通过入口校验（{len(errs)} 项）——不生成报告：", file=sys.stderr)
        for e in errs:
            print(f"  - {e}", file=sys.stderr)
        sys.exit(2)
    if warns:
        # 输出分层（P10）：报告清单里「同前缀多条」是常态，逐条刷屏无信息量——默认折叠为计数。
        if args.verbose:
            for w in warns:
                print(f"[WARN] {w}", file=sys.stderr)
        else:
            print(f"[WARN] 入口校验 {len(warns)} 条提示（同前缀多条属正常；如需逐条加 --verbose）",
                  file=sys.stderr)

    meta = {}
    if args.meta:
        try:
            with open(args.meta, encoding="utf-8") as fh:
                meta = json.load(fh)
        except (OSError, json.JSONDecodeError) as e:
            print(f"[ERROR] 元信息读取/解析失败：{e}", file=sys.stderr)
            sys.exit(2)
        if not isinstance(meta, dict):
            print("[ERROR] 元信息顶层须为对象（dict）", file=sys.stderr)
            sys.exit(2)

    # 声明未做 ＝ 未声明范围（与 SKIP 同纪律）：缺声明不阻断生成，但必须留痕
    if not meta.get("declaration"):
        print("[WARN] meta 未提供 declaration（复核声明）——封面「复核声明」节将缺省；"
              "凡非「全项」复核，按 ibd-doc-review delivery.md §八 须随报告声明核对范围。",
              file=sys.stderr)

    out, rows, cols = build(args, issues, meta)
    hi = sum(1 for r in rows if r["sev"] == "高")
    mid = sum(1 for r in rows if r["sev"] == "中")
    lo = sum(1 for r in rows if r["sev"] == "低")
    sheets = [n for n, _ in group_entries(rows, args.split)]

    if args.json:
        print(json.dumps({"tool": "review_report_to_xlsx", "out": out, "total": len(rows),
                          "split": args.split, "groups": sheets, "cols": [h for _, h in cols],
                          "declaration": bool(meta.get("declaration")),
                          "sev": {"高": hi, "中": mid, "低": lo}},
                         ensure_ascii=False))
        return 0
    # 输出分层（P10）：stdout 只出指标行 ＋ 路径，明细在工作簿里
    print(f"[OK] 复核报告已生成：{out}")
    print(f"     条目 {len(rows)}｜工作表 {len(sheets) + 2}（封面与汇总 ／ {' ／ '.join(sheets)} ／ 全部）")
    print(f"     严重度 高 {hi} ／ 中 {mid} ／ 低 {lo}｜分页轴 {args.split}｜列 {len(cols)} 列")
    return 0


if __name__ == "__main__":
    sys.exit(main())
