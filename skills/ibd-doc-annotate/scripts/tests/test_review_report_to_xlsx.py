#!/usr/bin/env python3
"""test_review_report_to_xlsx.py — 复核报告（Excel）生成器自测

三档样本（门禁/判据类改动的强制动作）：
  正常   多性质条目 ＋ 完整 meta → 封面页 ＋ 按性质分页 ＋ 全部页，列与编号符合规范
  回退   无 `nature` 字段 → 按 `type` 大类回退分页（不报错）
  边界   无 meta → 仍出封面（报告头结论由自动统计兜底）＋ 明细页；`--split sev` 换轴；坏清单 exit 2

覆盖（9 项）：多性质分表 ｜ nature 回退 type ｜ 无 meta 封面仍完整 ｜ `--split sev` ｜ 坏清单拦截
            ｜ **复核声明四栏渲染 ＋ 声明在先的动态编号** ｜ **只主标题居中·其余标题左对齐**
            ｜ **封面上下等宽（全行跨到 C 列）** ｜ **声明缺省时 WARN**

依赖：openpyxl（未安装则整类 SKIP——本功能按需装，不属本包基础依赖）。
运行：python scripts/tests/test_review_report_to_xlsx.py
"""
import json
import os
import subprocess
import sys
import tempfile
import unittest

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SKILL_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SCRIPT = os.path.join(SKILL_DIR, "scripts", "review_report_to_xlsx.py")

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAS_OPENPYXL = False


def _entry(no, nature, sev, code="A", **kw):
    it = {"author": "复核人甲", "code": code, "type": kw.pop("type", "数据·勾稽"), "sev": sev,
          "anchor": f"锚点片段{no}", "title": f"问题标题{no}", "desc": f"问题描述{no}",
          "advice": f"建议{no}"}
    if nature is not None:
        it["nature"] = nature
    it.update(kw)
    return it


def _run(td, issues, meta=None, extra=(), expect_rc=0):
    ip = os.path.join(td, "issues.json")
    with open(ip, "w", encoding="utf-8") as fh:
        json.dump(issues, fh, ensure_ascii=False)
    cmd = [sys.executable, SCRIPT, "--issues", ip, "--out", os.path.join(td, "报告.xlsx"), *extra]
    if meta is not None:
        mp = os.path.join(td, "meta.json")
        with open(mp, "w", encoding="utf-8") as fh:
            json.dump(meta, fh, ensure_ascii=False)
        cmd += ["--meta", mp]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if proc.returncode != expect_rc:
        raise AssertionError(f"rc={proc.returncode} 预期 {expect_rc}\nstdout:{proc.stdout}\nstderr:{proc.stderr}")
    return proc


def _load(td):
    return openpyxl.load_workbook(os.path.join(td, "报告.xlsx"))


def _cover_text(td, sheet="封面与汇总"):
    """封面所有非空单元格文字拼一行——用于「某段是否出现在封面」的断言。"""
    ws = openpyxl.load_workbook(os.path.join(td, "报告.xlsx"))[sheet]
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl 未安装")
class TestReviewReport(unittest.TestCase):

    def test_normal_multi_nature(self):
        """正常：多性质 → 封面 ＋ 5 张性质页 ＋ 全部；编号/列合规范。"""
        issues = [_entry(1, "①数据不一致", "高"), _entry(2, "①数据不一致", "中", code="B"),
                  _entry(3, "②表述不一致", "低"), _entry(4, "③自身逻辑", "中"),
                  _entry(5, "④格式规范", "低"), _entry(6, "⑤事实存疑", "中")]
        with tempfile.TemporaryDirectory() as td:
            _run(td, issues, meta={"title": "示例复核报告", "subject": "《回复》"})
            wb = _load(td)
            names = wb.sheetnames
            self.assertEqual(names[0], "封面与汇总")
            self.assertEqual(names[-1], "全部")
            self.assertEqual(len(names), 7, f"应为 封面+5性质+全部：{names}")
            for want in ("①数据不一致", "②表述不一致", "③自身逻辑", "④格式规范", "⑤事实存疑"):
                self.assertIn(want, names)
            # 分页顺序按**前导圈号**（标签文字项目可自定，圈号才稳定；如「③自身逻辑」/「③回复自身逻辑」）
            self.assertEqual([n for n in names if n[:1] in "①②③④⑤"],
                             ["①数据不一致", "②表述不一致", "③自身逻辑", "④格式规范", "⑤事实存疑"])
            ws = wb["全部"]
            header = [c.value for c in ws[1]]
            self.assertEqual(header[:5], ["编号", "严重度", "性质", "域", "行号"])
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 6)
            # 编号＝前缀＋前缀内序号（与注入器同源），严重度降序（高→低）
            self.assertEqual(rows[0][0], "A-01")
            self.assertEqual(rows[0][1], "高")
            self.assertEqual(rows[1][0], "B-01")
            self.assertEqual(rows[-1][1], "低")
            # 性质页只装本性质条目
            self.assertEqual(len(list(wb["①数据不一致"].iter_rows(min_row=2, values_only=True))), 2)

    def test_no_nature_falls_back_to_type(self):
        """回退：无 nature → 按 type 大类分页，不报错。"""
        issues = [_entry(1, None, "高", type="数据·口径"), _entry(2, None, "中", type="表述·歧义"),
                  _entry(3, None, "低", type="合规·一致性")]
        with tempfile.TemporaryDirectory() as td:
            _run(td, issues)
            names = _load(td).sheetnames
            for want in ("数据", "表述", "合规"):
                self.assertIn(want, names)
            self.assertNotIn("封面与汇总", names[1:])

    def test_no_meta_cover_still_complete(self):
        """边界：无 meta → 封面仍有报告头与自动结论统计；明细齐全。"""
        issues = [_entry(1, "①数据不一致", "高"), _entry(2, "②表述不一致", "中")]
        with tempfile.TemporaryDirectory() as td:
            proc = _run(td, issues, extra=("--json",))
            payload = json.loads(proc.stdout.strip().splitlines()[-1])
            self.assertEqual(payload["total"], 2)
            self.assertEqual(payload["sev"], {"高": 1, "中": 1, "低": 0})
            wb = _load(td)
            cover = wb["封面与汇总"]
            text = "\n".join(str(c.value) for row in cover.iter_rows() for c in row if c.value)
            self.assertIn("复核结论", text)
            self.assertIn("共确认 2 条问题", text)
            self.assertIn("按严重度", text)

    def test_declaration_rendered_on_cover(self):
        """正常：meta.declaration 四栏 → 封面「一、复核声明」节；声明在先，后续小节顺延编号。"""
        issues = [_entry(1, "①数据不一致", "高")]
        decl = {"depth": "L2 标准",
                "checked": ["格式规范：序号/错别字/标点/表格"],
                "not_checked": ["会计专业判断 —— 本次不出处理意见"],
                "special": ["S3 敏感议题：只标注、不出修改建议"],
                "gates": "交付件核验（F 域）10 项"}
        with tempfile.TemporaryDirectory() as td:
            proc = _run(td, issues,
                        meta={"title": "示例", "declaration": decl, "conclusion_notes": ["两条主线。"]},
                        extra=("--json",))
            self.assertTrue(json.loads(proc.stdout.strip().splitlines()[-1])["declaration"])
            text = _cover_text(td)
            for frag in ("一、复核声明", "复核深度：L2 标准",
                         "▍检查（本次实际要跑的项）", "▍不查（点名 ＋ 原因）",
                         "▍特别专项（触发即生效）", "▍交付前将跑",
                         "格式规范：序号/错别字/标点/表格", "会计专业判断 —— 本次不出处理意见",
                         "S3 敏感议题：只标注、不出修改建议", "交付件核验（F 域）10 项"):
                self.assertIn(frag, text)
            # 声明在先 ⇒ 结论为「二、」、统计为「三、」（编号按渲染顺序自增）
            self.assertIn("二、结论", text)
            self.assertIn("三、统计", text)
            self.assertNotIn("一、结论", text)

    def test_declaration_missing_warns(self):
        """边界：无 declaration → 不阻断生成，但 stderr 留痕（声明未做＝未声明范围）。"""
        issues = [_entry(1, "①数据不一致", "高")]
        with tempfile.TemporaryDirectory() as td:
            proc = _run(td, issues)
            self.assertIn("declaration", proc.stderr)
            self.assertNotIn("复核声明", _cover_text(td))

    def test_only_main_title_centered(self):
        """正常：**只第一行（主标题）居中**；副标题与小节标题一律左对齐（用户裁定）。"""
        issues = [_entry(1, "①数据不一致", "高")]
        with tempfile.TemporaryDirectory() as td:
            _run(td, issues, meta={"title": "居中示例报告", "declaration": {"checked": ["x"]}})
            wb = _load(td)
            ws = wb["封面与汇总"]
            merges = {str(r) for r in ws.merged_cells.ranges}
            title = ws.cell(1, 1)
            self.assertEqual(title.value, "居中示例报告")
            self.assertEqual(title.alignment.horizontal, "center")
            self.assertIn("A1:C1", merges)
            # 副标题（第 2 行）左对齐、不居中
            self.assertNotEqual(ws.cell(2, 1).alignment.horizontal, "center")
            # 小节标题：取「报告头」所在行，须左对齐
            row_of = {str(ws.cell(i, 1).value): i for i in range(1, ws.max_row + 1)}
            rr = row_of["报告头"]
            self.assertNotEqual(ws.cell(rr, 1).alignment.horizontal, "center")
            # **所有粗体大标题里，只有第 1 行居中**（表格表头居中属正常，不计）
            heads = [i for i in range(1, ws.max_row + 1)
                     if ws.cell(i, 1).font.bold and (ws.cell(i, 1).font.size or 0) >= 12]
            centered_heads = [i for i in heads if ws.cell(i, 1).alignment.horizontal == "center"]
            self.assertEqual(centered_heads, [1], f"只应主标题居中，实为：{centered_heads}")

    def test_cover_rows_uniform_width(self):
        """正常：封面**上下等宽**——所有有内容的行一律跨/铺到 C 列（用户裁定：统一 ABC 合并）。"""
        issues = [_entry(1, "①数据不一致", "高", code="A"), _entry(2, "②表述不一致", "中", code="B")]
        meta = {"title": "等宽示例", "conclusion_notes": ["结论行。"],
                "scope": "问题1–2。", "approach": ["① 基准单一；"],
                "declaration": {"depth": "L2", "checked": ["格式规范"]},
                "domain_scope": {"A": "问题1", "B": "问题2"},
                "verified": ["数据一致"], "quality_evidence": ["锚点全命中"], "actions": ["待决策"]}
        with tempfile.TemporaryDirectory() as td:
            _run(td, issues, meta=meta)
            ws = _load(td)["封面与汇总"]
            # 每行右边界＝合并区或带内容的最后一列的较大者
            edges = {}
            for rng in ws.merged_cells.ranges:
                for rr in range(rng.min_row, rng.max_row + 1):
                    edges[rr] = max(edges.get(rr, 0), rng.max_col)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None or c.border.right.style is not None:
                        edges[c.row] = max(edges.get(c.row, 0), c.column)
            bad = {r: e for r, e in edges.items() if e != 3}
            self.assertEqual(bad, {}, f"以下行未铺到 C 列（上下不等宽）：{bad}")
            # 三列列宽均已显式设定（不再有默认宽度的第三列）
            for col in ("A", "B", "C"):
                self.assertIsNotNone(ws.column_dimensions[col].width)

    def test_split_sev_axis(self):
        """边界：--split sev 换轴 → 三张严重度页。"""
        issues = [_entry(1, "①数据不一致", "高"), _entry(2, "①数据不一致", "中"), _entry(3, "②表述不一致", "低")]
        with tempfile.TemporaryDirectory() as td:
            _run(td, issues, extra=("--split", "sev"))
            names = _load(td).sheetnames
            self.assertEqual([n for n in names if n in ("高", "中", "低")], ["高", "中", "低"])

    def test_bad_issues_blocked(self):
        """阴性：缺 anchor 的坏清单 → exit 2、不出文件。"""
        with tempfile.TemporaryDirectory() as td:
            bad = [{"code": "A", "type": "数据·勾稽", "sev": "高", "title": "t", "desc": "d", "advice": "a"}]
            _run(td, bad, expect_rc=2)
            self.assertFalse(os.path.exists(os.path.join(td, "报告.xlsx")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
