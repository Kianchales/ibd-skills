#!/usr/bin/env python3
"""
test_annotate_docx.py — annotate_docx.py 注入器修复自测（T1 保留性 / T2 编号合规 / T3 总览只产 md）

背景（2026-09-29 注入器修复批次）：
  1. T1 段落重建原为「保留 pPr、清空其余内容流」——静默删除同段他人批注锚点、
     吞掉 w:tab 兄弟 run；且 comments.xml 被整体覆写、新批注 id 从 0 起与残留批注冲突。
     修复后：非文本内容原位保留、既有批注内容合并保留、新批注 id 接续既有最大 id。
  2. T2 同前缀超 99 条曾产出 3 位序号（X-103），与门禁 LABEL_PAT（前缀 [A-Z]{1,2}、
     序号恰 2 位）冲突；修复后顺延双字母分段（J-99 → JA-01…），序号恒 2 位。
  3. T3 总览**恒只产 md**（中间件、不交付用户）——2026-10-07 撤除 Word 版链路与
     `--overview-docx` 开关；人读报告改由 review_report_to_xlsx.py 出 Excel 复核报告。

覆盖：
  TestNumbering（纯函数，被测 = issue_numbering）：边界（1/99/100/131）、131 条同前缀全部合规、
    双字母前缀溢出报错、分段容量穷尽报错
  TestPreserveAndMerge（端到端）：① 纯 w:tab 兄弟 run 注入后仍在、段落文本零改动；
    ② 二次批注：既有 range 保留、comments.xml 合并（既有批注内容不丢）、id 接续、
    cs/ce/ref 对数与 comments 条数一致
  TestOverviewFlag（端到端）：总览恒只产 md、不产 docx；已撤除的 --overview-docx 须报错

依赖：python-docx（本包 docx 链路的声明依赖）——未安装则整类 SKIP。
运行：python scripts/tests/test_annotate_docx.py
"""
import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SKILL_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SCRIPT = os.path.join(SKILL_DIR, "scripts", "annotate_docx.py")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
# 门禁镜像：ibd-doc-review check_annotations.py 的 LABEL_PAT（前缀 1-2 位字母、序号恰 2 位）
LABEL_PAT = re.compile(r"^【([A-Z]{1,2})-(\d{2})｜")

try:
    from docx import Document
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False


def _make_fixture(path):
    """构建测试 docx：单段四 run——'AAAA' + 纯 w:tab run + '锚点文本在这里' + 'BBBB'。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("AAAA")
    p.add_run("锚点文本在这里")
    p.add_run("BBBB")
    # 在第二个 run 前插一个纯 w:tab run（r2 元素的 addprevious）
    runs = p._element.findall(f"{{{W}}}r")
    tab_run = OxmlElement("w:r")
    tab_run.append(OxmlElement("w:tab"))
    runs[1].addprevious(tab_run)
    doc.save(path)


def _issue(anchor, code="L", title="测试标题"):
    return {"author": "复核人甲", "code": code, "type": "表述不一致", "sev": "低",
            "anchor": anchor, "title": title, "desc": "问题描述", "advice": "修改建议"}


def _run_inject(src, issues, out, extra=()):
    issues_path = src + ".issues.json"
    with open(issues_path, "w", encoding="utf-8") as fh:
        json.dump(issues, fh, ensure_ascii=False)
    cmd = [sys.executable, SCRIPT, "--docx", src, "--issues", issues_path, "--out", out, *extra]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if proc.returncode != 0:
        raise AssertionError(f"注入失败 rc={proc.returncode}\nstdout: {proc.stdout}\nstderr: {proc.stderr}")
    return proc


def _read_part(path, part):
    with zipfile.ZipFile(path) as z:
        return z.read(part)


def _cs_ce_ref_counts(document_xml):
    root = ET.fromstring(document_xml)
    tags = {"cs": f"{{{W}}}commentRangeStart", "ce": f"{{{W}}}commentRangeEnd",
            "ref": f"{{{W}}}commentReference"}
    return {k: len(list(root.iter(t))) for k, t in tags.items()}


def _comments(path):
    root = ET.fromstring(_read_part(path, "word/comments.xml"))
    out = []
    for c in root.findall(f"{{{W}}}comment"):
        text = "".join(t.text or "" for t in c.iter(f"{{{W}}}t"))
        out.append((c.get(f"{{{W}}}id"), c.get(f"{{{W}}}author"), text))
    return out


class TestNumbering(unittest.TestCase):
    """T2：编号分段顺延（纯函数，被测 = issue_numbering；**零三方依赖，恒运行**）。

    2026-10-07 起编号规则抽至 `issue_numbering.py`，本类不再依赖 python-docx——
    故摘掉原与 docx 链路共用的 `skipUnless`（抽模块后该跳过条件已不成立）。
    """

    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, os.path.join(SKILL_DIR, "scripts"))
        # 编号规则的单一事实源已抽至 issue_numbering.py（2026-10-07）——
        # annotate_docx / annotate_pdf / review_report_to_xlsx 三处共用，勿再各自抄一份。
        import issue_numbering
        cls._full_label = staticmethod(issue_numbering.full_label)

    def test_boundaries(self):
        f = self._full_label
        self.assertEqual(f("J", 1), "J-01")
        self.assertEqual(f("J", 99), "J-99")
        self.assertEqual(f("J", 100), "JA-01")   # 第 100 条顺延双字母分段
        self.assertEqual(f("J", 131), "JA-32")   # 事故实测场景（131 条语言复核）
        self.assertEqual(f("J", 2673), "JZ-99")  # 容量上界
        self.assertEqual(f("AB", 99), "AB-99")   # 双字母前缀 99 条内正常

    def test_all_labels_match_gate_pattern(self):
        """131 条同前缀的编号全部过门禁镜像（前缀 [A-Z]{1,2}、序号恰 2 位）。"""
        labels = [self._full_label("J", n) for n in range(1, 132)]
        for lab in labels:
            self.assertRegex(f"【{lab}｜类型｜严重度】", LABEL_PAT)

    def test_two_letter_prefix_overflow_raises(self):
        with self.assertRaises(ValueError):
            self._full_label("AB", 100)

    def test_capacity_exhausted_raises(self):
        with self.assertRaises(ValueError):
            self._full_label("J", 2674)


@unittest.skipUnless(HAS_DOCX, "python-docx 未安装")
class TestPreserveAndMerge(unittest.TestCase):
    """T1：原位保留（tab run / 既有 range）＋ comments.xml 合并 ＋ id 接续（端到端）。"""

    def test_tab_run_and_second_pass(self):
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "fixture.docx")
            _make_fixture(src)

            # —— 第一批注入（锚点在 tab run 之后的 run 内）——
            out1 = os.path.join(td, "一批.docx")
            _run_inject(src, [_issue("锚点文本在这里", code="L", title="第一批标题")], out1)

            # tab run 存活（修复前会被段落重建吞掉）
            self.assertIn(b"<w:tab/>", _read_part(out1, "word/document.xml"))
            # 段落文本零改动（p.text 口径，w:tab 渲染为 \t）
            doc = Document(out1)
            self.assertEqual(doc.paragraphs[0].text, "AAAA\t锚点文本在这里BBBB")
            # 门禁对账：1 条批注 ↔ cs/ce/ref 各 1
            self.assertEqual(_cs_ce_ref_counts(_read_part(out1, "word/document.xml")),
                             {"cs": 1, "ce": 1, "ref": 1})
            self.assertEqual(len(_comments(out1)), 1)

            # —— 第二批注入（在第一批产物上二次批注，锚点覆盖段首 run）——
            out2 = os.path.join(td, "二批.docx")
            _run_inject(out1, [_issue("AAAA", code="M", title="第二批标题")], out2)

            doc_xml = _read_part(out2, "word/document.xml")
            # 既有 range 保留（修复前清空内容流会删掉第一批的 cs/ce/ref）
            self.assertEqual(_cs_ce_ref_counts(doc_xml), {"cs": 2, "ce": 2, "ref": 2})
            # tab run 仍在、文本仍零改动
            self.assertIn(b"<w:tab/>", doc_xml)
            self.assertEqual(Document(out2).paragraphs[0].text, "AAAA\t锚点文本在这里BBBB")
            # comments.xml 合并：2 条、id 接续 {0,1}、既有批注内容不丢
            cmts = _comments(out2)
            self.assertEqual(len(cmts), 2)
            self.assertEqual({c[0] for c in cmts}, {"0", "1"})
            self.assertTrue(any("第一批标题" in c[2] for c in cmts), "既有批注内容被覆写丢失")
            self.assertTrue(any("第二批标题" in c[2] for c in cmts))
            # 标签全部合规（门禁镜像）
            for _, _, text in cmts:
                self.assertRegex(text, LABEL_PAT)


@unittest.skipUnless(HAS_DOCX, "python-docx 未安装")
class TestOverviewFlag(unittest.TestCase):
    """T3（2026-10-07 起）：总览**恒只产 md**（中间件，不交付用户）；Word 版链路与
    `--overview-docx` 开关已撤除——旧调用须被显式拒绝，不得静默失效。"""

    def _inject(self, td, tag, extra):
        src = os.path.join(td, f"src_{tag}.docx")
        _make_fixture(src)
        out = os.path.join(td, f"out_{tag}.docx")
        _run_inject(src, [_issue("锚点文本在这里")], out, extra=extra)
        base = os.path.splitext(out)[0]
        return os.path.exists(base + "_批注总览.md"), os.path.exists(base + "_批注总览.docx")

    def test_only_md_never_docx(self):
        """阳性 + 阴性：总览 md 必产、总览 docx 恒不产。"""
        with tempfile.TemporaryDirectory() as td:
            md, docx = self._inject(td, "only_md", extra=())
            self.assertTrue(md, "总览 md 应恒产出（中间件）")
            self.assertFalse(docx, "总览 Word 版已撤除，不得产出")

    def test_removed_flag_errors(self):
        """边界：已撤除的 --overview-docx 须被 argcparse 拒绝（rc≠0），防旧调用静默失效。"""
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "src_flag.docx")
            _make_fixture(src)
            out = os.path.join(td, "out_flag.docx")
            issues_path = src + ".issues.json"
            with open(issues_path, "w", encoding="utf-8") as fh:
                json.dump([_issue("锚点文本在这里")], fh, ensure_ascii=False)
            cmd = [sys.executable, SCRIPT, "--docx", src, "--issues", issues_path,
                   "--out", out, "--overview-docx"]
            proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
            self.assertNotEqual(proc.returncode, 0, "已撤除的 --overview-docx 不应被接受")


if __name__ == "__main__":
    unittest.main(verbosity=2)
